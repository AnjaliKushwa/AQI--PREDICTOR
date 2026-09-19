from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation()

# Path to the generated images
title_bg = r"C:\Users\purva\.gemini\antigravity\brain\d87b9082-0000-442a-ba93-957cf89ed121\title_slide_bg_1778731477105.png"
ml_concept = r"C:\Users\purva\.gemini\antigravity\brain\d87b9082-0000-442a-ba93-957cf89ed121\ml_concept_1778731698225.png"
dashboard_mockup = r"C:\Users\purva\.gemini\antigravity\brain\d87b9082-0000-442a-ba93-957cf89ed121\dashboard_mockup_1778731722054.png"

# Slide 1: Title Slide with background image
blank_slide_layout = prs.slide_layouts[6]
slide1 = prs.slides.add_slide(blank_slide_layout)
try:
    slide1.shapes.add_picture(title_bg, 0, 0, width=prs.slide_width, height=prs.slide_height)
except:
    pass

# Add text box over image
txBox = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
tf = txBox.text_frame
p = tf.add_paragraph()
p.text = "AQI Predictor Pro"
p.font.size = Pt(60)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

p2 = tf.add_paragraph()
p2.text = "Empowering Citizens with Actionable Air Quality Insights"
p2.font.size = Pt(28)
p2.font.color.rgb = RGBColor(200, 255, 200)

# Slide 2: The Problem
slide_layout_content = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(slide_layout_content)
title_shape = slide2.shapes.title
body_shape = slide2.placeholders[1]
title_shape.text = "The Silent Crisis"
tf2 = body_shape.text_frame
tf2.text = "Millions are affected by air pollution, yet environmental data remains difficult to understand."
p = tf2.add_paragraph()
p.text = "Technical metrics (PM2.5, NO2, AQI numbers) fail to answer basic human questions: 'Is it safe to go outside today?'"
p.level = 0
p = tf2.add_paragraph()
p.text = "Existing platforms are built for scientists, not for everyday citizens."
p.level = 0

# Slide 3: The Solution with ML diagram
slide3 = prs.slides.add_slide(slide_layout_content)
title_shape = slide3.shapes.title
body_shape = slide3.placeholders[1]
title_shape.text = "Introducing AQI Predictor Pro"
tf3 = body_shape.text_frame
tf3.text = "Bridges the gap between raw scientific data and everyday health."
p = tf3.add_paragraph()
p.text = "Uses Machine Learning to predict accurate Air Quality Index (AQI) based on 6 key pollutants."
p.level = 0

# Insert ML image
try:
    slide3.shapes.add_picture(ml_concept, Inches(5), Inches(2.5), width=Inches(4.5))
except:
    pass

# Slide 4: Key Features (Dashboard)
slide4 = prs.slides.add_slide(prs.slide_layouts[5]) # Title only
title_shape = slide4.shapes.title
title_shape.text = "Visualizing the Impact"
try:
    # Insert Dashboard Mockup
    slide4.shapes.add_picture(dashboard_mockup, Inches(1.5), Inches(1.5), width=Inches(7))
except:
    pass

# Slide 5: Features Text
slide5 = prs.slides.add_slide(slide_layout_content)
title_shape = slide5.shapes.title
body_shape = slide5.placeholders[1]
title_shape.text = "Key Features"
tf5 = body_shape.text_frame
tf5.text = "ML-Powered Prediction: Evaluates PM2.5, PM10, NO2, SO2, CO, and O3."
p = tf5.add_paragraph()
p.text = "Plain-Language Advisory: Symptom trackers and time-of-day guidance."
p.level = 0
p = tf5.add_paragraph()
p.text = "Group-Specific Alerts: Tailored advice for children, elderly, and asthma patients."
p.level = 0
p = tf5.add_paragraph()
p.text = "City Heat Sink Map: Mapping of temperature, humidity, and pollution hotspots."
p.level = 0

# Slide 6: Technology Stack
slide6 = prs.slides.add_slide(slide_layout_content)
title_shape = slide6.shapes.title
body_shape = slide6.placeholders[1]
title_shape.text = "Technology Stack"
tf6 = body_shape.text_frame
tf6.text = "Frontend/UI: Streamlit (Python)"
p = tf6.add_paragraph()
p.text = "Machine Learning: Scikit-Learn"
p.level = 0
p = tf6.add_paragraph()
p.text = "Data Visualization: Plotly Graph Objects"
p.level = 0
p = tf6.add_paragraph()
p.text = "Styling: Custom CSS for a premium feel"
p.level = 0

prs.save("AQI_Project_Presentation.pptx")
print("Presentation successfully created at AQI_Project_Presentation.pptx")
